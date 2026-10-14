"""SQLite access layer: connect, inspect schema, run SQL, load the dataset."""
import csv
import datetime
import re
import sqlite3
import time
from pathlib import Path
from typing import Any

from app.config import DB_PATH, SALES_CSV_PATH


def get_connection() -> sqlite3.Connection:
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    # Some mounted/networked filesystems (e.g. a synced project folder) don't
    # support the POSIX locking SQLite's default rollback journal needs, and
    # fail with "disk I/O error". journal_mode=MEMORY avoids touching a
    # separate journal file on disk and works everywhere. Fine for a
    # single-process college project; revisit if you need multi-process
    # writers or crash-safe durability.
    conn.execute("PRAGMA journal_mode=MEMORY")
    return conn


SALES_TABLE = "sales"

# Generated joins can be accidentally unconstrained (a missing ON clause is a
# cross product — 50k x 50k rows), so cap how long any one query may run.
QUERY_TIMEOUT_S = 15


def _quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def list_tables() -> list[str]:
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ).fetchall()
        return [r["name"] for r in rows]
    finally:
        conn.close()


def list_uploaded_tables() -> list[str]:
    """Every table that came from a user upload — i.e. everything but `sales`."""
    return [t for t in list_tables() if t != SALES_TABLE]


def get_schema(tables: list[str] | None = None, with_samples: bool = False) -> str:
    """Return a compact text description of the given tables + their columns
    (every table if none given).

    This is what gets fed into the Ollama prompt so the model knows what it's
    allowed to query, and it's handy for debugging too. `with_samples` adds a
    few example values after each column, which is what lets the model tell
    that `orders.customer_id` and `customers.id` hold the same kind of value
    when it has to pick a join key.
    """
    names = tables if tables else list_tables()
    conn = get_connection()
    try:
        lines = []
        for tname in names:
            cols = conn.execute(f"PRAGMA table_info({_quote(tname)})").fetchall()
            parts = []
            for c in cols:
                desc = f"{c['name']} {c['type']}"
                if with_samples:
                    vals = conn.execute(
                        f"SELECT DISTINCT {_quote(c['name'])} FROM {_quote(tname)} "
                        f"WHERE {_quote(c['name'])} IS NOT NULL LIMIT 3"
                    ).fetchall()
                    if vals:
                        desc += " [e.g. " + ", ".join(str(v[0])[:20] for v in vals) + "]"
                parts.append(desc)
            lines.append(f"{tname}({', '.join(parts)})")
        return "\n".join(lines)
    finally:
        conn.close()


def get_relationships(tables: list[str]) -> list[dict]:
    """How the given tables probably relate, from three signals: `<x>_id`
    columns pointing at a table called `<x>` ("fk"), non-id columns that share
    a name ("shared"), and tables with identical columns, which can be stacked
    with UNION ALL ("stackable"). Shared by the join hints given to the
    model, the upload health card and the auto-drawn ER diagram.
    """
    conn = get_connection()
    try:
        cols = {
            t: [c["name"] for c in conn.execute(f"PRAGMA table_info({_quote(t)})").fetchall()]
            for t in tables
        }
    finally:
        conn.close()

    found: list[dict] = []

    by_name: dict[str, list[str]] = {}
    for t, names in cols.items():
        for n in names:
            by_name.setdefault(n, []).append(t)
    for n, ts in by_name.items():
        # Each table's own `id` is its primary key, not a link to another table.
        if len(ts) > 1 and n != "id":
            found.append({"kind": "shared", "column": n, "tables": ts})

    for t, names in cols.items():
        for n in names:
            if not n.endswith("_id"):
                continue
            stem = n[: -len("_id")]
            for other in tables:
                if other != t and other in (stem, stem + "s", stem + "es") and "id" in cols[other]:
                    found.append({"kind": "fk", "from_table": t, "from_col": n, "to_table": other, "to_col": "id"})

    for i, a in enumerate(tables):
        for b in tables[i + 1:]:
            if set(cols[a]) == set(cols[b]):
                found.append({"kind": "stackable", "tables": [a, b]})

    return found


def get_join_hints(tables: list[str]) -> str:
    """The relationships above as text for the model's prompt, so it doesn't
    have to guess a join key from column names alone. "" when nothing links
    the tables."""
    lines = []
    for r in get_relationships(tables):
        if r["kind"] == "shared":
            lines.append(f"- column `{r['column']}` appears in: {', '.join(r['tables'])}")
        elif r["kind"] == "fk":
            lines.append(f"- `{r['from_table']}.{r['from_col']}` likely references `{r['to_table']}.{r['to_col']}`")
        else:
            a, b = r["tables"]
            lines.append(
                f"- `{a}` and `{b}` have identical columns, so they *could* be stacked with UNION ALL — "
                "but only if the question asks to combine them or refers to both; otherwise use just the table it names"
            )
    return "\n".join(lines)


def execute_sql(sql: str, read_only: bool = False) -> tuple[list[str], list[tuple[Any, ...]]]:
    """Run a single SQL statement. Returns (column_names, rows).

    For SELECT this returns the result set. For INSERT/UPDATE/DELETE it
    commits and returns an empty result set with rowcount reported via
    columns=['rows_affected']. `read_only` makes SQLite itself refuse any
    write — used for SQL a person typed, where "it starts with SELECT" is
    not a strong enough guarantee on its own.
    """
    conn = get_connection()
    if read_only:
        conn.execute("PRAGMA query_only = ON")
    deadline = time.monotonic() + QUERY_TIMEOUT_S
    # Called every ~100k SQLite VM steps; a non-zero return aborts the query.
    conn.set_progress_handler(lambda: 1 if time.monotonic() > deadline else 0, 100_000)
    try:
        cur = conn.execute(sql)
        if sql.strip().lower().startswith("select"):
            columns = [d[0] for d in cur.description] if cur.description else []
            rows = cur.fetchall()
            return columns, [tuple(r) for r in rows]
        conn.commit()
        return ["rows_affected"], [(cur.rowcount,)]
    except sqlite3.OperationalError as exc:
        if "interrupted" in str(exc):
            raise TimeoutError(
                f"Query ran longer than {QUERY_TIMEOUT_S}s and was stopped — "
                "if it joins tables, it may be missing a join condition."
            ) from exc
        raise
    finally:
        conn.close()


_SALES_COLUMNS = [
    "product",
    "brand",
    "product_code",
    "product_spec",
    "price",
    "inward_date",
    "dispatch_date",
    "quantity_sold",
    "customer_name",
    "customer_location",
    "region",
    "core_spec",
    "processor_spec",
    "ram",
    "rom",
    "ssd",
]


def _clean_row(row: dict) -> tuple:
    def s(key: str) -> str | None:
        val = (row.get(key) or "").strip()
        return val or None

    def n(key: str) -> int | None:
        val = (row.get(key) or "").strip()
        try:
            return int(float(val)) if val else None
        except ValueError:
            return None

    return (
        s("Product"),
        s("Brand"),
        s("Product Code"),
        s("Product Specification"),
        n("Price"),
        s("Inward Date"),
        s("Dispatch Date"),
        n("Quantity Sold"),
        s("Customer Name"),
        s("Customer Location"),
        s("Region"),
        s("Core Specification"),
        s("Processor Specification"),
        s("RAM"),
        s("ROM"),
        s("SSD"),
    )


def load_sales_data() -> None:
    """Create the `sales` table and load it from SALES_CSV_PATH, if empty.

    Idempotent: does nothing if the table already has rows, so restarts
    don't re-import 50k rows every time.
    """
    conn = get_connection()
    try:
        numeric_cols = {"price": "INTEGER", "quantity_sold": "INTEGER"}
        cols_sql = ",\n                ".join(
            f"{c} {numeric_cols.get(c, 'TEXT')}" for c in _SALES_COLUMNS
        )
        conn.execute(
            f"""
            CREATE TABLE IF NOT EXISTS sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                {cols_sql}
            )
            """
        )
        count = conn.execute("SELECT COUNT(*) FROM sales").fetchone()[0]
        if count > 0:
            return

        csv_path = Path(SALES_CSV_PATH)
        if not csv_path.exists():
            return  # nothing to import; table stays empty

        placeholders = ", ".join("?" for _ in _SALES_COLUMNS)
        insert_sql = f"INSERT INTO sales ({', '.join(_SALES_COLUMNS)}) VALUES ({placeholders})"

        with csv_path.open(newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            batch = []
            for row in reader:
                batch.append(_clean_row(row))
                if len(batch) >= 2000:
                    conn.executemany(insert_sql, batch)
                    batch.clear()
            if batch:
                conn.executemany(insert_sql, batch)
        conn.commit()
    finally:
        conn.close()


# Table names that would need quoting everywhere if a CSV were named after
# them ("order.csv", "group.csv") — cheaper to just rename the table.
_SQL_KEYWORDS = {
    "all", "as", "asc", "by", "case", "check", "column", "cross", "default",
    "desc", "end", "from", "group", "if", "in", "index", "is", "join", "key",
    "left", "like", "limit", "natural", "not", "null", "on", "or", "order",
    "primary", "references", "right", "select", "table", "to", "transaction",
    "trigger", "union", "using", "values", "view", "where", "with",
}


def _sanitize_identifier(name: str, fallback: str) -> str:
    cleaned = re.sub(r"\W+", "_", name.strip()).strip("_").lower()
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"{fallback}_{cleaned}" if cleaned else fallback
    return cleaned


def _table_name(stem: str) -> str:
    name = _sanitize_identifier(stem, "dataset")
    if name == SALES_TABLE:
        name = "sales_upload"
    if name in _SQL_KEYWORDS:
        name = f"{name}_data"
    return name


def _dedupe_columns(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for col in columns:
        if col not in seen:
            seen[col] = 0
            result.append(col)
        else:
            seen[col] += 1
            result.append(f"{col}_{seen[col]}")
    return result


def _infer_column_types(rows: list[list[str]], n_cols: int) -> list[str]:
    types = []
    for i in range(n_cols):
        values = [r[i].strip() for r in rows[:500] if r[i].strip()]
        if not values:
            types.append("TEXT")
        elif all(re.fullmatch(r"-?\d+", v) for v in values):
            types.append("INTEGER")
        elif all(re.fullmatch(r"-?\d+(\.\d+)?", v) for v in values):
            types.append("REAL")  # includes a mix like 10 and 12.5
        else:
            types.append("TEXT")
    return types


def _coerce(val: str, col_type: str):
    val = val.strip()
    if not val:
        return None
    try:
        if col_type == "INTEGER":
            return int(val)
        if col_type == "REAL":
            return float(val)
    except ValueError:
        return None
    return val


def _load_table(table: str, raw_columns: list[str], rows: list[list[str]]) -> dict:
    """Create `table` from a header row and string-valued rows, replacing any
    table already called that. Shared by every upload format, so a CSV and an
    Excel sheet end up as identical tables.

    Column names are sanitized to valid SQL identifiers and types inferred by
    sampling. There is deliberately no synthetic `id` column: a file's own
    `id` column is the natural join key, so it must keep its name.
    """
    columns = _dedupe_columns(
        [_sanitize_identifier(c, f"col{i}") for i, c in enumerate(raw_columns)]
    )
    width = len(columns)
    rows = [(list(r) + [""] * width)[:width] for r in rows]  # ragged rows -> rectangular
    col_types = _infer_column_types(rows, width)

    qtable = _quote(table)
    qcols = [_quote(c) for c in columns]

    conn = get_connection()
    try:
        replaced = table in {r["name"] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()}
        conn.execute(f"DROP TABLE IF EXISTS {qtable}")
        cols_sql = ",\n                ".join(f"{q} {t}" for q, t in zip(qcols, col_types))
        conn.execute(f"CREATE TABLE {qtable} (\n                {cols_sql}\n            )")

        placeholders = ", ".join("?" for _ in columns)
        insert_sql = f"INSERT INTO {qtable} ({', '.join(qcols)}) VALUES ({placeholders})"

        batch = []
        for row in rows:
            batch.append(tuple(_coerce(v, t) for v, t in zip(row, col_types)))
            if len(batch) >= 2000:
                conn.executemany(insert_sql, batch)
                batch.clear()
        if batch:
            conn.executemany(insert_sql, batch)
        conn.commit()
    finally:
        conn.close()

    return {"table": table, "columns": columns, "row_count": len(rows), "replaced": replaced}


def _cell_to_str(v: Any) -> str:
    """An Excel cell as the text a CSV export would have contained, so the
    same type inference applies (5.0 -> "5", midnight datetimes -> a date)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _import_csv(csv_path: Path, original_filename: str) -> list[dict]:
    with csv_path.open(newline="", encoding="utf-8-sig") as f:
        grid = [r for r in csv.reader(f) if any(c.strip() for c in r)]
    if not grid:
        raise ValueError("CSV has no header row")
    info = _load_table(_table_name(Path(original_filename).stem), grid[0], grid[1:])
    return [{**info, "filename": original_filename, "sheet": None}]


def _import_xlsx(xlsx_path: Path, original_filename: str) -> list[dict]:
    from openpyxl import load_workbook  # only needed for Excel uploads

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheets: list[tuple[str, list[str], list[list[str]]]] = []
    try:
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()  # some writers record a wrong used-range
            except AttributeError:
                pass
            grid = [list(r) for r in ws.iter_rows(values_only=True)]
            grid = [r for r in grid if any(c is not None and str(c).strip() for c in r)]
            if not grid:
                continue
            width = 1 + max(
                i for r in grid for i, c in enumerate(r) if c is not None and str(c).strip()
            )
            text = [[_cell_to_str(c) for c in (list(r) + [None] * width)[:width]] for r in grid]
            sheets.append((ws.title, text[0], text[1:]))
    finally:
        wb.close()

    if not sheets:
        raise ValueError("The workbook has no data")

    stem = Path(original_filename).stem
    results = []
    for title, header, body in sheets:
        # One sheet: named after the file. Several: "<file>_<sheet>", so two
        # workbooks that both have a "Sheet1" can't overwrite each other.
        name = _table_name(stem if len(sheets) == 1 else f"{stem}_{title}")
        info = _load_table(name, header, body)
        results.append({**info, "filename": original_filename, "sheet": title})
    return results


def import_uploaded_file(path: Path, original_filename: str) -> list[dict]:
    """Load an uploaded CSV or Excel workbook into the database — one table
    per CSV, or one per non-empty sheet of a workbook. Tables are named after
    the file ("customers.csv" -> `customers`), any number can sit side by side
    so a question can join/merge them, and uploading a name that already exists
    replaces that table. They have no hand-written rule-based parser support
    (unlike `sales`), so they're only ever queried via the Ollama fallback,
    which reads the schema straight from the DB.
    """
    ext = Path(original_filename).suffix.lower()
    if ext == ".csv":
        return _import_csv(path, original_filename)
    if ext in (".xlsx", ".xlsm"):
        return _import_xlsx(path, original_filename)
    raise ValueError("Please upload a .csv or .xlsx file.")


def describe_uploaded_tables() -> list[dict]:
    """Name, columns and row count of every uploaded table — lets the UI
    rebuild its dataset list after a page reload."""
    conn = get_connection()
    try:
        out = []
        for t in list_uploaded_tables():
            cols = [c["name"] for c in conn.execute(f"PRAGMA table_info({_quote(t)})").fetchall()]
            count = conn.execute(f"SELECT COUNT(*) FROM {_quote(t)}").fetchone()[0]
            out.append({"table": t, "columns": cols, "row_count": count})
        return out
    finally:
        conn.close()


def drop_uploaded_table(table: str) -> None:
    if table not in list_uploaded_tables():
        raise ValueError(f"No uploaded dataset named {table!r}.")
    conn = get_connection()
    try:
        conn.execute(f"DROP TABLE {_quote(table)}")
        conn.commit()
    finally:
        conn.close()
